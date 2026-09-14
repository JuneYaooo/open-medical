"""Structured Office layout, used only inside the bounded document runtime."""
import re
import tempfile
import subprocess
from pathlib import Path


def color(value, fallback="233247"):
    value = str(value or fallback).lstrip("#")
    if not re.fullmatch(r"[0-9a-fA-F]{6}", value):
        raise ValueError("Colors must be six hexadecimal digits")
    return value


def number(value, default, minimum, maximum):
    result = float(default if value is None else value)
    if not minimum <= result <= maximum:
        raise ValueError("Layout dimension is outside supported bounds")
    return result


def blocks(content):
    result = content.get("blocks", [{"type": "paragraph", "text": str(p)} for p in content.get("paragraphs", [])])
    if len(result) > 500:
        raise ValueError("At most 500 document blocks")
    return result


def table_rows(block):
    rows = block.get("rows", [])
    if not rows or len(rows) > 1000 or not rows[0] or len(rows[0]) > 30 or any(len(r) != len(rows[0]) for r in rows):
        raise ValueError("Tables require a rectangular grid, up to 1000 rows and 30 columns")
    return rows


def create_docx(path, content, safe_path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    theme = content.get("theme", {})
    font = theme.get("font", "Noto Sans CJK SC")
    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(number(theme.get("fontSize"), 11, 8, 36))
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    normal.paragraph_format.space_after = Pt(8)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(number(theme.get("margin"), 0.8, 0.25, 2))
    section.left_margin = section.right_margin = section.top_margin
    if content.get("header"):
        section.header.paragraphs[0].text = str(content["header"])
    footer = section.footer.paragraphs[0]
    footer.text = str(content.get("footer", ""))
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    if content.get("title"):
        doc.add_heading(content["title"], 0)
    for block in blocks(content):
        kind = block.get("type", "paragraph")
        if kind == "pageBreak":
            doc.add_page_break()
        elif kind == "heading":
            doc.add_heading(str(block.get("text", "")), int(number(block.get("level"), 1, 1, 6)))
        elif kind == "table":
            rows = table_rows(block)
            table = doc.add_table(rows=0, cols=len(rows[0])); table.style = "Light Shading Accent 1"
            for index, values in enumerate(rows):
                row = table.add_row()
                for cell, value in zip(row.cells, values):
                    cell.text = str(value)
                if index == 0:
                    repeat = OxmlElement("w:tblHeader"); row._tr.get_or_add_trPr().append(repeat)
                    for cell in row.cells:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
        elif kind == "image":
            doc.add_picture(str(safe_path(block["path"])), width=Inches(number(block.get("width"), 5.5, 0.1, 7)))
            if block.get("caption"):
                doc.add_paragraph(str(block["caption"]), "Caption")
        elif kind in ("paragraph", "bullet", "numbered"):
            paragraph = doc.add_paragraph(style={"bullet": "List Bullet", "numbered": "List Number"}.get(kind))
            for item in block.get("runs", [{"text": block.get("text", "")} ]):
                run = paragraph.add_run(str(item.get("text", "")))
                run.bold = item.get("bold", False); run.italic = item.get("italic", False)
                if item.get("color"):
                    run.font.color.rgb = RGBColor.from_string(color(item["color"]))
            paragraph.alignment = {"left": 0, "center": 1, "right": 2, "justify": 3}.get(block.get("align"), 0)
            paragraph.paragraph_format.keep_with_next = bool(block.get("keepWithNext", False))
        else:
            raise ValueError("Unknown document block: " + kind)
    doc.save(path)


def create_xlsx(path, content, safe_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.drawing.image import Image
    book = Workbook()
    sheets = content.get("sheets", [{"name": "Sheet1", "rows": []}])
    if not 1 <= len(sheets) <= 30:
        raise ValueError("Workbooks support 1 to 30 sheets")
    for index, spec in enumerate(sheets):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = spec["name"]
        rows = spec.get("rows", [])
        if len(rows) > 10000 or any(len(row) > 100 for row in rows):
            raise ValueError("Sheet exceeds 10000 rows or 100 columns")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = spec.get("freeze", "A2")
        for row in sheet:
            for cell in row:
                cell.font = Font(name="Noto Sans CJK SC", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if rows and spec.get("header", True):
            for cell in sheet[1]:
                cell.font = Font(name="Noto Sans CJK SC", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=color(spec.get("accent")))
            sheet.row_dimensions[1].height = 26
        for column, width in spec.get("widths", {}).items():
            if not re.fullmatch(r"[A-Z]{1,3}", column):
                raise ValueError("Column width keys must be column letters")
            sheet.column_dimensions[column].width = number(width, 18, 2, 100)
        for address, fmt in spec.get("formats", {}).items():
            for row in sheet[address] if ":" in address else [(sheet[address],)]:
                for cell in row:
                    cell.number_format = str(fmt)
        for merged in spec.get("merges", []):
            sheet.merge_cells(merged)
        if spec.get("table") and len(rows) > 1:
            if any(not isinstance(v, str) for v in rows[0]):
                raise ValueError("Table headers must be strings")
            table = Table(displayName="DataTable" + str(index + 1), ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        for spec_chart in spec.get("charts", [])[:20]:
            chart_type = spec_chart.get("type", "bar")
            if chart_type not in ("bar", "line", "pie"):
                raise ValueError("Chart type must be bar, line or pie")
            chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[chart_type]()
            chart.title = spec_chart.get("title", "")
            chart.add_data(Reference(sheet, range_string=f"'{sheet.title}'!{spec_chart['data']}"), titles_from_data=True)
            if spec_chart.get("categories"):
                chart.set_categories(Reference(sheet, range_string=f"'{sheet.title}'!{spec_chart['categories']}"))
            sheet.add_chart(chart, spec_chart.get("anchor", "E2"))
        for picture in spec.get("images", [])[:20]:
            sheet.add_image(Image(str(safe_path(picture["path"]))), picture.get("anchor", "A1"))
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = spec.get("orientation", "landscape")
        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
        sheet.oddFooter.center.text = "Page &P / &N"
    book.save(path)


def create_pptx(path, content, safe_path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    deck = Presentation(); deck.slide_width = Inches(13.333); deck.slide_height = Inches(7.5)
    slides = content.get("slides", [])
    if not 1 <= len(slides) <= 100:
        raise ValueError("Presentations support 1 to 100 slides")
    def textbox(slide, text, x, y, width, height, size=20, bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        shape.text_frame.word_wrap = True
        for i, line in enumerate(str(text).split("\n")):
            p = shape.text_frame.paragraphs[0] if i == 0 else shape.text_frame.add_paragraph()
            p.text = line; p.font.name = "Noto Sans CJK SC"; p.font.size = Pt(size); p.font.bold = bold
        return shape
    for number_slide, spec in enumerate(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(color(spec.get("background"), "FFFFFF"))
        textbox(slide, spec.get("title", ""), .65, .4, 12, .8, 30, True)
        if spec.get("paragraphs"):
            textbox(slide, "\n".join(spec["paragraphs"]), .7, 1.5, 11.8, 5, 22)
        for block in spec.get("elements", [])[:100]:
            x = number(block.get("x"), .7, 0, 13); y = number(block.get("y"), 1.5, 0, 7.3)
            width = number(block.get("width"), 5, .1, 13.333 - x); height = number(block.get("height"), 3, .1, 7.5 - y)
            kind = block.get("type", "text")
            if kind == "text":
                textbox(slide, block.get("text", ""), x, y, width, height, number(block.get("fontSize"), 20, 8, 80), block.get("bold", False))
            elif kind == "image":
                slide.shapes.add_picture(str(safe_path(block["path"])), Inches(x), Inches(y), width=Inches(width), height=Inches(height))
            elif kind == "table":
                rows = table_rows(block)
                table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(width), Inches(height)).table
                for row_index, row in enumerate(rows):
                    for col_index, value in enumerate(row):
                        cell = table.cell(row_index, col_index); cell.text = str(value)
                        for p in cell.text_frame.paragraphs:
                            p.font.size = Pt(14); p.font.name = "Noto Sans CJK SC"
            elif kind == "chart":
                data = CategoryChartData(); data.categories = block["categories"]
                for series in block["series"]:
                    data.add_series(series["name"], series["values"])
                chart_kind = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE}.get(block.get("chartType", "bar"))
                if chart_kind is None: raise ValueError("Unknown chart type")
                slide.shapes.add_chart(chart_kind, Inches(x), Inches(y), Inches(width), Inches(height), data)
            else:
                raise ValueError("Unknown slide element: " + kind)
        textbox(slide, str(number_slide + 1), 12, 7, .6, .3, 10)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(spec["notes"])
    deck.save(path)


def create_pdf(path, content, safe_path):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from xml.sax.saxutils import escape
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    style = ParagraphStyle("body", fontName="STSong-Light", fontSize=11, leading=17, wordWrap="CJK", spaceAfter=8)
    heading = ParagraphStyle("heading", parent=style, fontSize=20, leading=26, spaceAfter=16)
    items = []
    def paragraph(value, selected=style):
        return Paragraph(escape(str(value)).replace("\n", "<br/>"), selected)
    if content.get("title"): items.append(paragraph(content["title"], heading))
    for block in blocks(content):
        kind = block.get("type", "paragraph")
        if kind == "pageBreak": items.append(PageBreak())
        elif kind == "table":
            rows = table_rows(block)
            table = Table([[paragraph(v) for v in row] for row in rows], colWidths=[480 / len(rows[0])] * len(rows[0]), repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E4EAF0")), ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#BCC6D0")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
            items.extend([table, Spacer(1, 12)])
        elif kind == "image":
            image = Image(str(safe_path(block["path"])))
            scale = min(480 / image.imageWidth, 550 / image.imageHeight, 1)
            image.drawWidth = image.imageWidth * scale; image.drawHeight = image.imageHeight * scale
            items.append(image)
            if block.get("caption"): items.append(paragraph(block["caption"]))
        elif kind in ("paragraph", "heading", "bullet", "numbered"):
            text = block.get("text", "")
            if "runs" in block: text = "".join(str(run.get("text", "")) for run in block["runs"])
            items.append(paragraph(("• " if kind == "bullet" else "") + str(text), heading if kind == "heading" else style))
        else: raise ValueError("Unknown PDF block: " + kind)
    def page(canvas, doc):
        canvas.saveState(); canvas.setFont("STSong-Light", 9)
        canvas.drawString(50, doc.pagesize[1] - 30, str(content.get("header", "")))
        canvas.drawString(50, 28, str(content.get("footer", "")))
        canvas.drawRightString(doc.pagesize[0] - 50, 28, str(doc.page)); canvas.restoreState()
    SimpleDocTemplate(str(path), leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50).build(items, onFirstPage=page, onLaterPages=page)


def render_document(path, safe_path):
    if path.suffix.lower() not in (".docx", ".xlsx", ".pptx", ".pdf"):
        raise ValueError("Only Office and PDF rendering is supported")
    destination = safe_path(path.stem + "-preview")
    if destination.exists():
        raise ValueError("Preview directory already exists")
    destination.mkdir()
    with tempfile.TemporaryDirectory() as profile:
        if path.suffix.lower() == ".pdf":
            pdf = path
        else:
            subprocess.run(["libreoffice", "-env:UserInstallation=" + Path(profile).as_uri(), "--headless", "--convert-to", "pdf", "--outdir", str(destination), str(path)], check=True, timeout=90, capture_output=True)
            pdf = destination / (path.stem + ".pdf")
        if not pdf.is_file(): raise ValueError("Renderer did not produce PDF")
        subprocess.run(["pdftoppm", "-png", "-scale-to", "1200", "-f", "1", "-l", "20", str(pdf), str(destination / "page")], check=True, timeout=90, capture_output=True)
    return {"pages": [str(p.relative_to(Path.cwd())) for p in sorted(destination.glob("page-*.png"))], "pdf": str(pdf.relative_to(Path.cwd())), "pageLimit": 20}
